#!/usr/bin/env python3
"""SPDR sector holdings capture/publish regressions -- fixture .xlsx
workbook + fake HTTP layer only. No network call is ever made by this
file, and the real SSGA endpoint is never contacted."""
from __future__ import annotations

import datetime as dt
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "spdr_sector_holdings", ROOT / "collectors" / "spdr_sector_holdings.py"
)
M = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(M)

NOW = dt.datetime(2026, 9, 15, 6, 0, 0, tzinfo=dt.timezone.utc)


def fixture_workbook(rows: list[tuple], header=("Ticker", "Name", "Weight (%)")) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def fixture_workbook_with_preamble(
    preamble_rows: list, rows: list[tuple], header=("Ticker", "Name", "Weight (%)")
) -> bytes:
    """Like fixture_workbook, but with extra header rows (e.g. an "as of"
    date cell, a fund title) BEFORE the real column-header row -- exactly
    the shape parse_holdings_workbook/parse_holdings_as_of_date must
    tolerate."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for preamble_row in preamble_rows:
        ws.append(list(preamble_row))
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class SectorUniverseTests(unittest.TestCase):
    def test_universe_matches_repo_config(self):
        import json
        contract = json.loads((ROOT / "config" / "free_market_data_contract.json").read_text())
        reference = set(contract["alpaca"]["sector_reference_symbols"]) - {"SMH"}
        self.assertEqual(set(M.SECTOR_ETFS), reference)
        self.assertEqual(len(M.SECTOR_ETFS), 11)


class ParsingTests(unittest.TestCase):
    def test_parses_symbol_and_weight_and_skips_cash(self):
        raw = fixture_workbook([
            ("NVDA", "NVIDIA CORP", 8.5),
            ("AAPL", "APPLE INC", 4.2),
            ("CASH", "CASH", 0.1),
        ])
        holdings = M.parse_holdings_workbook(raw)
        self.assertEqual({h["symbol"] for h in holdings}, {"NVDA", "AAPL"})

    def test_tolerates_alternate_header_names(self):
        raw = fixture_workbook(
            [("MSFT", "MICROSOFT", "6.10%")],
            header=("Identifier", "Security Description", "Weight"),
        )
        holdings = M.parse_holdings_workbook(raw)
        self.assertEqual(holdings[0]["symbol"], "MSFT")
        self.assertAlmostEqual(holdings[0]["weight_pct"], 6.10)

    def test_leading_junk_rows_before_header_are_skipped(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["SPDR Select Sector Fund - XLK", None, None])
        ws.append(["As of 2026-09-14", None, None])
        ws.append(["Ticker", "Name", "Weight (%)"])
        ws.append(["NVDA", "NVIDIA", 8.5])
        buf = io.BytesIO()
        wb.save(buf)
        holdings = M.parse_holdings_workbook(buf.getvalue())
        self.assertEqual(holdings[0]["symbol"], "NVDA")

    def test_unreadable_bytes_fail_closed(self):
        with self.assertRaisesRegex(M.SpdrSectorHoldingsError, "HOLDINGS_WORKBOOK_UNREADABLE"):
            M.parse_holdings_workbook(b"not an xlsx file")

    def test_no_header_found_fails_closed(self):
        raw = fixture_workbook([], header=("Foo", "Bar", "Baz"))
        with self.assertRaisesRegex(M.SpdrSectorHoldingsError, "HOLDINGS_HEADER_NOT_FOUND"):
            M.parse_holdings_workbook(raw)


class WeightBucketTests(unittest.TestCase):
    def test_buckets(self):
        self.assertEqual(M.weight_bucket(0.5), "LT_1PCT")
        self.assertEqual(M.weight_bucket(1.0), "GE_1PCT_LT_5PCT")
        self.assertEqual(M.weight_bucket(4.99), "GE_1PCT_LT_5PCT")
        self.assertEqual(M.weight_bucket(5.0), "GE_5PCT")


class HoldingsAsOfDateTests(unittest.TestCase):
    """collectors/spdr_sector_holdings.py::parse_holdings_as_of_date --
    the real SSGA header wording/format is UNVERIFIED (see module
    docstring); this exercises several plausible spellings tolerantly."""

    def test_slash_date_with_are_as_of_phrasing(self):
        raw = fixture_workbook_with_preamble(
            [["Holdings are as of 09/12/2026", None, None]], [("NVDA", "NVIDIA", 8.5)],
        )
        self.assertEqual(M.parse_holdings_as_of_date(raw), "2026-09-12")

    def test_iso_date_with_colon_phrasing(self):
        raw = fixture_workbook_with_preamble(
            [["SPDR XLK", None, None], ["As Of: 2026-09-12", None, None]],
            [("NVDA", "NVIDIA", 8.5)],
        )
        self.assertEqual(M.parse_holdings_as_of_date(raw), "2026-09-12")

    def test_month_name_date(self):
        raw = fixture_workbook_with_preamble(
            [["As Of September 12, 2026"]], [("NVDA", "NVIDIA", 8.5)],
        )
        self.assertEqual(M.parse_holdings_as_of_date(raw), "2026-09-12")

    def test_dd_mon_yyyy_date(self):
        raw = fixture_workbook_with_preamble(
            [["as-of 12-Sep-2026"]], [("NVDA", "NVIDIA", 8.5)],
        )
        self.assertEqual(M.parse_holdings_as_of_date(raw), "2026-09-12")

    def test_no_as_of_phrase_returns_none(self):
        raw = fixture_workbook_with_preamble(
            [["SPDR Select Sector Fund - XLK"]], [("NVDA", "NVIDIA", 8.5)],
        )
        self.assertIsNone(M.parse_holdings_as_of_date(raw))

    def test_as_of_phrase_without_a_parseable_date_returns_none(self):
        raw = fixture_workbook_with_preamble(
            [["Holdings are as of the most recent close"]], [("NVDA", "NVIDIA", 8.5)],
        )
        self.assertIsNone(M.parse_holdings_as_of_date(raw))

    def test_unreadable_workbook_returns_none_not_an_exception(self):
        self.assertIsNone(M.parse_holdings_as_of_date(b"not an xlsx file"))

    def test_decoy_date_before_the_as_of_phrase_in_the_same_cell_is_ignored(self):
        # PR #767 review: an unrelated earlier date in the same cell must
        # never be picked up as the "as of" date -- only a date
        # POSITIONALLY anchored right after the phrase counts.
        raw = fixture_workbook_with_preamble(
            [["Fund inception 01/01/2001. Holdings as of 09/12/2026"]],
            [("NVDA", "NVIDIA", 8.5)],
        )
        self.assertEqual(M.parse_holdings_as_of_date(raw), "2026-09-12")

    def test_conflicting_as_of_dates_across_rows_is_unknown(self):
        # An earlier disclaimer row with its own "as of <date>" must never
        # silently "win" over the real holdings row -- two distinct
        # anchored candidates means UNKNOWN, not a guess.
        raw = fixture_workbook_with_preamble(
            [["Data as of 01/01/2001 (report basis)"], ["Holdings are as of 09/12/2026"]],
            [("NVDA", "NVIDIA", 8.5)],
        )
        self.assertIsNone(M.parse_holdings_as_of_date(raw))

    def test_repeated_identical_as_of_date_across_rows_is_not_a_conflict(self):
        raw = fixture_workbook_with_preamble(
            [["See disclosures. As of 09/12/2026"], ["Holdings are as of 09/12/2026"]],
            [("NVDA", "NVIDIA", 8.5)],
        )
        self.assertEqual(M.parse_holdings_as_of_date(raw), "2026-09-12")


class FetchTests(unittest.TestCase):
    def test_fetch_holdings_uses_the_documented_url_template(self):
        calls = []
        def fake_getter(url, headers=None):
            calls.append(url)
            return fixture_workbook([("NVDA", "NVIDIA", 8.5)])
        M.fetch_holdings("XLK", getter=fake_getter)
        self.assertEqual(calls[0], M.holdings_url("XLK"))
        self.assertIn("holdings-daily-us-en-xlk.xlsx", calls[0])

    def test_fetch_rejects_ticker_outside_universe(self):
        with self.assertRaisesRegex(M.SpdrSectorHoldingsError, "HOLDINGS_TICKER_NOT_IN_UNIVERSE"):
            M.fetch_holdings("SMH", getter=lambda url, headers=None: b"")


class CaptureAndPublishTests(unittest.TestCase):
    def test_authority_is_all_false(self):
        raw = fixture_workbook([("NVDA", "NVIDIA", 8.5)])
        bundle = M.build_capture(NOW, "XLK", raw)
        for key, value in bundle["capture"]["authority"].items():
            if key.endswith("_authorized"):
                self.assertFalse(value, key)

    def test_derived_mapping_ranked_by_weight_descending(self):
        raw = fixture_workbook([
            ("AAPL", "APPLE", 4.2),
            ("NVDA", "NVIDIA", 8.5),
            ("AVGO", "BROADCOM", 4.2),
        ])
        bundle = M.build_capture(NOW, "XLK", raw)
        mapping = bundle["capture"]["mapping"]
        self.assertEqual(mapping[0]["symbol"], "NVDA")
        self.assertEqual(mapping[0]["weight_rank"], 1)
        self.assertEqual(mapping[0]["weight_bucket"], "GE_5PCT")

    def test_raw_workbook_bytes_are_never_stored_in_the_capture(self):
        raw = fixture_workbook([("NVDA", "NVIDIA", 8.5)])
        bundle = M.build_capture(NOW, "XLK", raw)
        serialized = bundle["capture_bytes"]
        self.assertNotIn(raw, serialized)
        self.assertIn(M.sha256_bytes(raw), bundle["capture"]["raw_sha256"])
        self.assertNotIn("openpyxl", str(bundle["capture"]))  # sanity: no binary leakage

    def test_publish_is_idempotent_and_content_addressed(self):
        raw = fixture_workbook([("NVDA", "NVIDIA", 8.5)])
        bundle = M.build_capture(NOW, "XLK", raw)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = M.publish_capture(root, bundle)
            self.assertTrue(first["created"])
            second = M.publish_capture(root, bundle)
            self.assertFalse(second["created"])
            self.assertEqual(first["capture_id"], second["capture_id"])

    def test_append_only_collision_fails_closed(self):
        raw = fixture_workbook([("NVDA", "NVIDIA", 8.5)])
        bundle = M.build_capture(NOW, "XLK", raw)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            M.publish_capture(root, bundle)
            (root / bundle["capture_path"]).write_bytes(b"tampered")
            with self.assertRaisesRegex(M.SpdrSectorHoldingsError, "APPEND_ONLY_COLLISION"):
                M.publish_capture(root, bundle)

    def test_two_batches_sharing_an_evidence_day_do_not_collide(self):
        # The as-of date does not move between every pair of scheduled runs:
        # the 2026-09-16 and 2026-09-17 batches on main both carry as-of
        # 2026-09-15. Re-publishing the same batch content under one
        # evidence_day must be a no-op, not APPEND_ONLY_COLLISION, even though
        # captured_at_utc differs.
        raw = {ticker: fixture_workbook_with_preamble(
            [["Holdings are as of 09/15/2026", None, None]], [("NVDA", "NVIDIA", 8.5)],
        ) for ticker in M.SECTOR_ETFS}
        first = M.build_batch(NOW, raw)
        second = M.build_batch(NOW + dt.timedelta(hours=22), raw)
        self.assertEqual(first["manifest_path"], second["manifest_path"])
        self.assertNotEqual(first["manifest_bytes"], second["manifest_bytes"])
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            M.publish_batch(root, first)
            M.publish_batch(root, second)  # no raise
            stored = json.loads((root / first["manifest_path"]).read_bytes())
            self.assertEqual(stored["captured_at_utc"], first["manifest"]["captured_at_utc"])

    def test_a_different_batch_under_one_evidence_day_still_fails_closed(self):
        as_of = [["Holdings are as of 09/15/2026", None, None]]
        raw_a = {t: fixture_workbook_with_preamble(as_of, [("NVDA", "NVIDIA", 8.5)]) for t in M.SECTOR_ETFS}
        raw_b = {t: fixture_workbook_with_preamble(as_of, [("AMD", "AMD", 4.25)]) for t in M.SECTOR_ETFS}
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            M.publish_batch(root, M.build_batch(NOW, raw_a))
            with self.assertRaisesRegex(M.SpdrSectorHoldingsError, "APPEND_ONLY_COLLISION"):
                M.publish_batch(root, M.build_batch(NOW + dt.timedelta(hours=22), raw_b))

    def test_evidence_is_keyed_by_the_holdings_as_of_date(self):
        # Two fetches on one UTC day with different as-of dates must land in
        # different directories: GitHub fires this collector hours late, so
        # 2026-09-17 saw 00:03Z (as-of 09-15) and 22:10Z (a newer as-of) and
        # the second failed with APPEND_ONLY_COLLISION.
        raw_old = fixture_workbook_with_preamble(
            [["Holdings are as of 09/15/2026", None, None]], [("NVDA", "NVIDIA", 8.5)],
        )
        raw_new = fixture_workbook_with_preamble(
            [["Holdings are as of 09/16/2026", None, None]], [("NVDA", "NVIDIA", 9.0)],
        )
        early = M.build_capture(NOW, "XLK", raw_old)
        late = M.build_capture(NOW + dt.timedelta(hours=22), "XLK", raw_new)
        self.assertNotEqual(early["capture_path"], late["capture_path"])
        self.assertIn("2026-09-15", early["capture_path"])
        self.assertIn("2026-09-16", late["capture_path"])
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.assertTrue(M.publish_capture(root, early)["created"])
            self.assertTrue(M.publish_capture(root, late)["created"])

    def test_capture_day_stays_the_key_when_the_as_of_date_is_unknown(self):
        raw = fixture_workbook([("NVDA", "NVIDIA", 8.5)])
        bundle = M.build_capture(NOW, "XLK", raw)
        self.assertEqual(bundle["capture"]["holdings_as_of_date"], M.HOLDINGS_AS_OF_UNKNOWN)
        self.assertEqual(bundle["capture"]["evidence_day"], bundle["capture"]["capture_date_utc"])
        self.assertIn(bundle["capture"]["capture_date_utc"], bundle["capture_path"])

    def test_same_workbook_recapture_later_the_same_day_is_a_no_op(self):
        raw = fixture_workbook([("NVDA", "NVIDIA", 8.5)])
        first_bundle = M.build_capture(NOW, "XLK", raw)
        later = NOW + dt.timedelta(hours=6)
        second_bundle = M.build_capture(later, "XLK", raw)
        self.assertEqual(first_bundle["capture_path"], second_bundle["capture_path"])
        self.assertNotEqual(first_bundle["capture_bytes"], second_bundle["capture_bytes"])
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            M.publish_capture(root, first_bundle)
            result = M.publish_capture(root, second_bundle)
            self.assertFalse(result["created"])
            kept = json.loads((root / first_bundle["capture_path"]).read_bytes())
            self.assertEqual(kept["captured_at_utc"], first_bundle["capture"]["captured_at_utc"])

    def test_different_workbook_under_the_same_path_still_fails_closed(self):
        first_bundle = M.build_capture(NOW, "XLK", fixture_workbook([("NVDA", "NVIDIA", 8.5)]))
        second_bundle = M.build_capture(NOW, "XLK", fixture_workbook([("AMD", "AMD", 4.25)]))
        self.assertEqual(first_bundle["capture_path"], second_bundle["capture_path"])
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            M.publish_capture(root, first_bundle)
            with self.assertRaisesRegex(M.SpdrSectorHoldingsError, "APPEND_ONLY_COLLISION"):
                M.publish_capture(root, second_bundle)

    def test_path_traversal_is_rejected(self):
        raw = fixture_workbook([("NVDA", "NVIDIA", 8.5)])
        bundle = M.build_capture(NOW, "XLK", raw)
        bundle["capture_path"] = f"{M.EVIDENCE_ROOT}/derived/../../secret.json"
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(M.SpdrSectorHoldingsError, "EVIDENCE_PATH_INVALID"):
                M.publish_capture(Path(tmp), bundle)

    def test_ticker_outside_universe_rejected_at_build_time(self):
        raw = fixture_workbook([("NVDA", "NVIDIA", 8.5)])
        with self.assertRaisesRegex(M.SpdrSectorHoldingsError, "HOLDINGS_TICKER_NOT_IN_UNIVERSE"):
            M.build_capture(NOW, "SMH", raw)

    def test_holdings_as_of_date_is_recorded_when_present(self):
        raw = fixture_workbook_with_preamble(
            [["Holdings are as of 09/12/2026"]], [("NVDA", "NVIDIA", 8.5)],
        )
        bundle = M.build_capture(NOW, "XLK", raw)
        self.assertEqual(bundle["capture"]["holdings_as_of_date"], "2026-09-12")

    def test_holdings_as_of_date_is_unknown_not_a_guess_when_absent(self):
        raw = fixture_workbook([("NVDA", "NVIDIA", 8.5)])  # no preamble at all
        bundle = M.build_capture(NOW, "XLK", raw)
        self.assertEqual(bundle["capture"]["holdings_as_of_date"], M.HOLDINGS_AS_OF_UNKNOWN)
        # Never silently defaults to the capture date -- that would assume
        # same-day freshness, which this feature exists to stop doing.
        self.assertNotEqual(bundle["capture"]["holdings_as_of_date"], bundle["capture"]["capture_date_utc"])


class LatestPointerTests(unittest.TestCase):
    def test_pointer_reports_completeness(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            per_ticker = {"XLK": {"capture_id": "abc"}}
            path = M.write_latest_pointer(root, "2026-09-15", per_ticker)
            import json
            pointer = json.loads(path.read_text())
            self.assertFalse(pointer["complete"])
            per_ticker_all = {t: {"capture_id": "x"} for t in M.SECTOR_ETFS}
            M.write_latest_pointer(root, "2026-09-15", per_ticker_all)
            pointer_all = json.loads(path.read_text())
            self.assertTrue(pointer_all["complete"])


ALL_11_RAW = {ticker: fixture_workbook([("PLACEHOLDER", "PLACEHOLDER CO", 0.5)]) for ticker in M.SECTOR_ETFS}


def all_11_raw_with(overrides: dict[str, bytes]) -> dict[str, bytes]:
    raw = dict(ALL_11_RAW)
    raw.update(overrides)
    return raw


class ResolveCrossEtfWinnersTests(unittest.TestCase):
    def test_sole_holder(self):
        resolved = M.resolve_cross_etf_winners({"XLK": [{"symbol": "NVDA", "weight_pct": 8.5, "name": None}]})
        self.assertEqual(resolved["NVDA"], {
            "schema_version": M.RESOLVED_SYMBOL_SCHEMA_VERSION,
            "symbol": "NVDA", "primary_sector_etf": "XLK",
            "holder_etf_count": 1, "tie": False,
        })

    def test_exact_weight_multi_holder_picks_the_true_larger_weight(self):
        # A rank-1-in-fund holding is NOT always heavier than a rank-2 one
        # in a different fund -- this is exactly the bug PR #761 fixed.
        per_ticker_holdings = {
            "XLK": [{"symbol": "NVDA", "weight_pct": 6.0, "name": None},
                    {"symbol": "AAPL", "weight_pct": 4.0, "name": None}],
            "XLC": [{"symbol": "META", "weight_pct": 9.0, "name": None},
                    {"symbol": "NVDA", "weight_pct": 6.5, "name": None}],
        }
        resolved = M.resolve_cross_etf_winners(per_ticker_holdings)
        self.assertEqual(resolved["NVDA"]["primary_sector_etf"], "XLC")  # 6.5 > 6.0
        self.assertEqual(resolved["NVDA"]["holder_etf_count"], 2)
        self.assertFalse(resolved["NVDA"]["tie"])

    def test_exact_tie_sets_tie_flag_and_breaks_alphabetically(self):
        per_ticker_holdings = {
            "XLK": [{"symbol": "DUP", "weight_pct": 5.0, "name": None}],
            "XLC": [{"symbol": "DUP", "weight_pct": 5.0, "name": None}],
        }
        resolved = M.resolve_cross_etf_winners(per_ticker_holdings)
        self.assertTrue(resolved["DUP"]["tie"])
        self.assertEqual(resolved["DUP"]["primary_sector_etf"], "XLC")


class BuildBatchTests(unittest.TestCase):
    def test_complete_batch_writes_resolved_symbols(self):
        raw = all_11_raw_with({
            "XLK": fixture_workbook([("NVDA", "NVIDIA", 8.5)]),
            "XLC": fixture_workbook([("META", "META", 9.0)]),
        })
        batch = M.build_batch(NOW, raw)
        self.assertTrue(batch["batch_complete"])
        self.assertIsNotNone(batch["symbols_bytes"])
        self.assertEqual(batch["manifest"]["tickers_captured"], sorted(M.SECTOR_ETFS))

    def test_incomplete_batch_has_no_resolved_symbols_but_still_has_per_etf_captures(self):
        raw = {"XLK": fixture_workbook([("NVDA", "NVIDIA", 8.5)])}
        batch = M.build_batch(NOW, raw)
        self.assertFalse(batch["batch_complete"])
        self.assertIsNone(batch["symbols_bytes"])
        self.assertIsNone(batch["symbols_path"])
        self.assertEqual(list(batch["per_ticker_capture"]), ["XLK"])

    def test_empty_batch_is_incomplete_and_does_not_crash(self):
        batch = M.build_batch(NOW, {})
        self.assertFalse(batch["batch_complete"])
        self.assertEqual(batch["per_ticker_capture"], {})

    def test_ticker_outside_universe_rejected(self):
        with self.assertRaisesRegex(M.SpdrSectorHoldingsError, "HOLDINGS_TICKER_NOT_IN_UNIVERSE"):
            M.build_batch(NOW, {"SMH": fixture_workbook([("NVDA", "NVIDIA", 8.5)])})

    def test_batch_aggregate_as_of_date_when_every_etf_agrees(self):
        with_as_of = fixture_workbook_with_preamble(
            [["Holdings are as of 09/12/2026"]], [("PLACEHOLDER", "PLACEHOLDER CO", 0.5)],
        )
        raw = {ticker: with_as_of for ticker in M.SECTOR_ETFS}
        batch = M.build_batch(NOW, raw)
        self.assertEqual(batch["manifest"]["holdings_as_of_date"], "2026-09-12")
        self.assertTrue(all(d == "2026-09-12" for d in batch["manifest"]["holdings_as_of_dates"].values()))

    def test_batch_aggregate_is_unknown_on_disagreement(self):
        raw = all_11_raw_with({
            "XLK": fixture_workbook_with_preamble([["as of 09/12/2026"]], [("NVDA", "NVIDIA", 8.5)]),
            "XLC": fixture_workbook_with_preamble([["as of 09/11/2026"]], [("META", "META", 9.0)]),
        })
        batch = M.build_batch(NOW, raw)
        self.assertEqual(batch["manifest"]["holdings_as_of_dates"]["XLK"], "2026-09-12")
        self.assertEqual(batch["manifest"]["holdings_as_of_dates"]["XLC"], "2026-09-11")
        self.assertEqual(batch["manifest"]["holdings_as_of_date"], M.HOLDINGS_AS_OF_UNKNOWN)

    def test_batch_aggregate_is_unknown_when_any_etf_is_unknown(self):
        # ALL_11_RAW's default fixture has no "as of" preamble at all --
        # even one ETF's holdings_as_of_date is UNKNOWN, so a naive
        # "agreement among the known ones" reading must NOT be used.
        raw = all_11_raw_with({
            "XLK": fixture_workbook_with_preamble([["as of 09/12/2026"]], [("NVDA", "NVIDIA", 8.5)]),
        })
        batch = M.build_batch(NOW, raw)
        self.assertEqual(batch["manifest"]["holdings_as_of_date"], M.HOLDINGS_AS_OF_UNKNOWN)


class PublishBatchTests(unittest.TestCase):
    def test_publish_complete_batch_writes_manifest_and_symbols(self):
        raw = all_11_raw_with({"XLK": fixture_workbook([("NVDA", "NVIDIA", 8.5)])})
        batch = M.build_batch(NOW, raw)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            summary = M.publish_batch(root, batch)
            self.assertTrue(summary["batch_complete"])
            self.assertTrue((root / summary["manifest_path"]).exists())
            self.assertTrue((root / summary["symbols_path"]).exists())
            self.assertEqual(len(summary["per_ticker"]), 11)
            import json
            symbols = json.loads((root / summary["symbols_path"]).read_text())
            self.assertIn("NVDA", [s["symbol"] for s in symbols])

    def test_publish_incomplete_batch_writes_manifest_only(self):
        raw = {"XLK": fixture_workbook([("NVDA", "NVIDIA", 8.5)])}
        batch = M.build_batch(NOW, raw)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            summary = M.publish_batch(root, batch)
            self.assertFalse(summary["batch_complete"])
            self.assertIsNone(summary["symbols_path"])
            self.assertTrue((root / summary["manifest_path"]).exists())

    def test_publish_is_idempotent(self):
        raw = all_11_raw_with({"XLK": fixture_workbook([("NVDA", "NVIDIA", 8.5)])})
        batch = M.build_batch(NOW, raw)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            M.publish_batch(root, batch)
            M.publish_batch(root, batch)  # no error, no collision


class TickerAllowlistCliTests(unittest.TestCase):
    """collectors/spdr_sector_holdings.py main() -- see
    .github/workflows/spdr-sector-holdings.yml's env:/quoted-variable fix
    (PR #761 review item 2): the shell never substitutes the untrusted
    workflow_dispatch input directly into the script text, and this script
    is responsible for splitting + validating it against SECTOR_ETFS."""

    def test_invalid_ticker_in_argv_is_rejected_before_any_fetch(self):
        calls = []
        original_fetch = M.fetch_holdings
        def spy(ticker, *, getter=None):
            calls.append(ticker)
            return original_fetch(ticker, getter=lambda url, headers=None: fixture_workbook([("X", "X", 1.0)]))
        M.fetch_holdings = spy
        try:
            rc = M.main(["--tickers", "XLK NOT_A_TICKER"])
        finally:
            M.fetch_holdings = original_fetch
        self.assertEqual(rc, 2)
        self.assertEqual(calls, [])  # rejected before any network call

    def test_space_separated_string_is_split_correctly(self):
        seen = []
        original_fetch = M.fetch_holdings
        def spy(ticker, *, getter=None):
            seen.append(ticker)
            raise M.SpdrSectorHoldingsError("HOLDINGS_HTTP_UNREACHABLE")
        M.fetch_holdings = spy
        try:
            with tempfile.TemporaryDirectory() as tmp:
                original_root = M.ROOT
                M.ROOT = Path(tmp)
                try:
                    rc = M.main(["--tickers", "XLK XLF"])
                finally:
                    M.ROOT = original_root
        finally:
            M.fetch_holdings = original_fetch
        self.assertEqual(rc, 0)  # per-ticker fetch failure does not crash the run
        self.assertEqual(sorted(seen), ["XLF", "XLK"])


if __name__ == "__main__":
    unittest.main()
